from docx import Document


def export_review(
    review
):

    document = Document()

    document.add_heading(
        "Literature Review",
        level=1
    )

    document.add_paragraph(
        review
    )

    filename = (
        "literature_review.docx"
    )

    document.save(
        filename
    )

    return filename