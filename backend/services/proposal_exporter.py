from docx import Document


def export_proposal(
    proposal
):

    document = Document()

    document.add_heading(

        "Research Proposal",

        level=1
    )

    document.add_paragraph(
        proposal
    )

    filename = (
        "research_proposal.docx"
    )

    document.save(
        filename
    )

    return filename